# -*- coding: utf-8 -*-
"""Convert PAPER_DRAFT.md (headings, runs, tables, images, lists, code, quotes)
into a clean university-style .docx with embedded figures."""
import os
import re
import sys

import docx
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding="utf-8")

SRC = r"C:\Users\Рам\Desktop\BitNet_Research_Reports\PAPER_DRAFT.md"
OUT = r"C:\Users\Рам\Desktop\BitNet_Research_Reports\СТАТЬЯ_BitNet_агенты_черновик.docx"
BASE = os.path.dirname(SRC)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+?\*)")


def add_runs(par, text):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = par.add_run(piece[2:-2])
            r.bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.italic = True
        else:
            par.add_run(piece)


def flush_table(doc, rows):
    if not rows:
        return
    # rows: list of lists of cell strings; first row = header
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            txt = row[j] if j < len(row) else ""
            add_runs(par, txt)
            for r in par.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


def main():
    doc = Document()
    # page setup: A4, 2cm margins
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin = sec.bottom_margin = Mm(20)

    # base styles
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5
    for lvl, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12.5)):
        h = doc.styles[lvl]
        h.font.name = "Times New Roman"
        h.font.size = Pt(size)
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.font.bold = True

    lines = open(SRC, encoding="utf-8").read().splitlines()
    i = 0
    table_rows = []
    in_code = False
    code_buf = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(8)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                code_buf = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # table rows
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):  # skip separator
                table_rows.append(cells)
            i += 1
            continue
        else:
            flush_table(doc, table_rows)
            table_rows = []

        s = line.strip()
        if not s:
            i += 1
            continue
        if s == "---":
            i += 1
            continue

        # image
        m = re.match(r"!\[(.*?)\]\((.+?)\)", s)
        if m:
            path = os.path.join(BASE, m.group(2))
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Mm(155))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(m.group(1))
                r.italic = True
                r.font.size = Pt(10)
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue

        # bullet / numbered
        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            i += 1
            continue

        # blockquote
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(8)
            add_runs(p, s.lstrip("> ").strip())
            for r in p.runs:
                r.font.size = Pt(11)
                r.italic = True
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        add_runs(p, s)
        i += 1

    flush_table(doc, table_rows)
    doc.save(OUT)
    print("saved:", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
